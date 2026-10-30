#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
写字楼经营分析周报 · BI 自动取数脚本
====================================
从观远 BI（地址由环境变量 BI_BASE 提供）逐卡拉取数据，产出供报告流水线消费的标准 Excel。

用法:
    python fetch_bi_data.py [--date YYYY-MM-DD] [--out 输出路径.xlsx]
    --date 缺省 = 昨日（报告截止日）

参数默认值规则（用户 2026-10-16 确认，详见 数据源映射.md 一·B）:
    查询日期   = 报告截止日（昨日）
    本期开始   = 2026-01-01（固定）    本期结束 = 报告截止日
    开始时间   = 2026-01-01（固定）    结束时间 = 报告截止日

产出 Excel sheet（不追求与原模板同构，以能出报告为准）:
    A基础信息(出租+收入) / A出租率分月 / A租费收入分月 / A累计补充分析
    B权责1-N月累计分析（复杂报表 HTML 解析成二维表，保持页面行列布局）
    B收入支出明细 / B出租面积明细 / B固定填报（子卡 CSV，交叉校验用）
    C合同情况(本年累计) / C流失统计(本年累计)
    D办公业态竞品
    F储备转化分析 / F储备明细
    取数日志（每卡行数、警告；任一关键卡 0 行 → 留空+报警，严禁编造）

依赖: 仅标准库 + openpyxl；认证 token 从 guancli 配置读取（路径可用 GUANCLI_CONFIG 覆盖）。

环境变量（首次使用按 README「配置」设置）:
    BI_BASE           必填，观远 BI 地址（如 https://bi.example.com）
    GUANCLI_BIN       可选，guancli 可执行文件（默认 PATH 里的 guancli）
    GUANCLI_CONFIG    可选，guancli config.json 路径（默认按平台）
    BI_PROFILE        可选，guancli profile 名（默认 jinying）
    PROJECT_ROOT      可选，数据/输出根目录（默认当前目录）
    BI_COMPLEX_CDID   可选，权责分析表复杂报表卡 ID
    BI_COMPLEX_BODY_TPL 可选，复杂报表请求体模板路径（默认 scripts/exact_body_tpl.json）
    BI_CARDS_FILE     可选，卡片清单 JSON（默认内置 12 张卡；重建卡片时可覆盖）
    NODE_BIN          可选，node 可执行文件（默认 PATH 里的 node，用于认证刷新）
    NODE_PATH         可选，Node 模块搜索路径
    OA_URL / OA_SSO_URL / BI_ORIGIN / OA_LOGIN_ID / OA_PASSWORD  认证刷新用，见 bi_auth_refresh.js
"""

import argparse
import csv
import datetime as dt
import io
import json
import os
import re
import shutil
import subprocess
import sys
import urllib.request

# ---------------------------------------------------------------- 常量（全部可由环境变量覆盖）
def _default_guancli_config():
    if sys.platform == "darwin":
        return "~/Library/Application Support/guancli/config.json"
    if os.name == "nt":
        return os.path.join(os.environ.get("APPDATA", ""), "guancli", "config.json")
    return "~/.config/guancli/config.json"


GUANCLI = shutil.which(os.environ.get("GUANCLI_BIN") or "guancli") or (
    os.environ.get("GUANCLI_BIN") or "guancli"
)
PROFILE = os.environ.get("BI_PROFILE") or "jinying"
BASE = (os.environ.get("BI_BASE") or "").rstrip("/")
TOKEN_CFG = os.path.expanduser(
    os.environ.get("GUANCLI_CONFIG") or _default_guancli_config()
)
PROJECT = os.environ.get("PROJECT_ROOT") or os.getcwd()
SCRIPTS_DIR = os.path.dirname(os.path.abspath(__file__))

COMPLEX_CDID = os.environ.get("BI_COMPLEX_CDID") or "je5ae866d4baf480e9a82235"   # 权责分析表（复杂报表卡）
COMPLEX_BODY_TPL = os.environ.get("BI_COMPLEX_BODY_TPL") or os.path.join(
    SCRIPTS_DIR, "exact_body_tpl.json"
)

# 页面日期参数 dpId（数据源映射.md 一·B）
DP = {
    "查询日期": "j6984e01cb5944fbebaddc1d",
    "本期开始": "m9df5b3cdc99f4f69b9961b2",
    "本期结束": "c73247308343a42e1be5d9ec",
    "开始时间": "m6a3e11c4043b470fa5f2107",
    "结束时间": "u8ec55ee779a0405c9e89ff7",
}

# 卡片清单: (sheet 名, cdId, 动态参数名列表, limit, 是否关键卡)
CARDS = [
    ("A基础信息(出租+收入)",  "hb23b7b6c6c1b4c688596ab8", ["查询日期"], 500, True),
    ("A出租率分月",           "c5beeff97c49f4dcdba2b833", ["查询日期"], 2000, False),
    ("A租费收入分月",         "d0607c0ae6fcd4506b94a3d0", ["查询日期"], 2000, False),
    ("A累计补充分析",         "t126122f425db498b886f1e2", ["查询日期"], 500, False),
    ("B收入支出明细",         "na08d034de8c74f68a69ddfa", ["查询日期"], 500, True),
    ("B出租面积明细",         "w05b234f6effd41c2ae91eb7", ["查询日期"], 500, False),
    ("B固定填报",             "q09ca4e4837c941bc98c335a", ["查询日期"], 500, False),
    ("C合同情况(本年累计)",   "c2b2ccfdb42f24a79950c456", ["查询日期", "本期开始", "本期结束"], 500, True),
    ("C流失统计(本年累计)",   "c0c7f679510d44180817c9e8", ["开始时间", "结束时间"], 500, False),
    ("D办公业态竞品",         "k5425cb5956c9495ea1d85d7", [], 500, False),
    ("F储备转化分析",         "baa8cba33009645d9a309882", ["开始时间", "结束时间"], 1000, True),
    ("F储备明细",             "s4483b6dc5773484fb8c5c43", ["开始时间", "结束时间"], 2000, False),
]

# 可选：用 JSON 覆盖卡片清单（BI 侧重建卡片时无需改代码）
_cards_file = os.environ.get("BI_CARDS_FILE")
if _cards_file:
    with open(_cards_file, encoding="utf-8") as _fh:
        CARDS = [tuple(item) for item in json.load(_fh)]

log_lines = []
warnings = []


def log(msg):
    stamp = dt.datetime.now().strftime("%H:%M:%S")
    line = f"[{stamp}] {msg}"
    print(line, flush=True)
    log_lines.append(line)


def warn(msg):
    warnings.append(msg)
    log(f"⚠️ {msg}")


# ---------------------------------------------------------------- 认证
def get_token():
    if not os.path.exists(TOKEN_CFG):
        sys.exit(
            f"❌ 未找到 guancli 配置：{TOKEN_CFG}\n"
            "请先安装 guancli 并完成 BI 登录，或用 GUANCLI_CONFIG 指定配置文件路径"
        )
    def find(o):
        if isinstance(o, dict):
            for k, v in o.items():
                if k in ("token", "uIdToken", "uidToken") and isinstance(v, str) and len(v) > 50:
                    return v
                r = find(v)
                if r:
                    return r
        elif isinstance(o, list):
            for v in o:
                r = find(v)
                if r:
                    return r
        return None
    with open(TOKEN_CFG, encoding="utf-8") as f:
        tok = find(json.load(f))
    if not tok:
        sys.exit("❌ 未能从 guancli config.json 提取 uIdToken，请重新配置认证")
    return tok


# ---------------------------------------------------------------- 普通卡取数（guancli CSV）
def fetch_card_csv(cdid, params, limit):
    cmd = [GUANCLI, "card", "preview", cdid, "--profile", PROFILE,
           "--limit", str(limit), "-f", "csv"]
    for name, val in params.items():
        cmd += ["--dynamic-param", f"{name}={val}"]
    p = subprocess.run(cmd, capture_output=True, text=True, timeout=300)
    out = p.stdout.strip()
    if p.returncode != 0 or not out:
        err = (p.stderr or p.stdout or "").strip()[:300]
        raise RuntimeError(f"guancli 失败: {err}")
    rows = list(csv.reader(io.StringIO(out)))
    return rows


def rows_all_zero(rows):
    """数据行（除表头）的所有数值列全为 0/空 → 视为该日期分区无数据。
    注意：跳过首列（维度列，如项目名），否则项目名文本会误判为「有数据」；
    非数值列仅认「明细类」实质文本（长度>1）为有数据。"""
    if len(rows) <= 1:
        return True
    for r in rows[1:]:
        for v in r[1:]:
            s = str(v).replace(",", "").strip()
            if not s:
                continue
            try:
                if float(s) != 0:
                    return False
            except ValueError:
                # 非数值列：有实质文本（如合同明细）也算有数据
                if len(s) > 1:
                    return False
    return True


def fetch_with_fallback(sheet, cdid, pnames, limit, param_vals, end_date, zero_fallback=False):
    """带分区回退的取数。「查询日期」是上游 dt 分区选择器，各源表刷新节奏不同
    （合同表比财务表慢 1 天），分区无数据时自动逐日回退，最多 7 天；
    「本期结束/结束时间」是期间边界，按用户规则固定=报告截止日，不参与回退。
    回退触发：0 行，或（zero_fallback=True 时）数值列全 0。
    返回 (rows, used_partition_date, note)。"""
    if "查询日期" not in pnames:
        params = {n: param_vals[n] for n in pnames}
        rows = fetch_card_csv(cdid, params, limit)
        return rows, None, ""
    d = dt.date.fromisoformat(end_date)
    rows, cur = [], end_date
    for back in range(8):
        cur = (d - dt.timedelta(days=back)).isoformat()
        params = {n: (cur if n == "查询日期" else param_vals[n]) for n in pnames}
        rows = fetch_card_csv(cdid, params, limit)
        empty = len(rows) <= 1
        allzero = zero_fallback and rows_all_zero(rows)
        if not empty and not allzero:
            if back:
                warn(f"{sheet}: 查询日期={end_date} 分区无数据，已回退至 {cur}（期间边界不变）")
            return rows, cur, (f"查询日期分区回退 {back} 天（{end_date} → {cur}）" if back else "")
        if back == 0:
            log(f"{sheet}: 查询日期={cur} 分区无数据，尝试回退…")
    warn(f"{sheet}: 查询日期连续回退 7 天仍无数据（{end_date} 起），上游分区疑似断更 → sheet 留空+报警")
    return rows, cur, f"回退 7 天仍无数据（最后尝试 {cur}）"


# ---------------------------------------------------------------- 复杂报表卡取数（API + HTML 解析）
def fetch_complex_report(token, query_date, period_start, period_end):
    """POST /api/card/{cdId}/data，dynamicParams[].defaultValue 传日期（逆向成果，见映射文档三·B）"""
    with open(COMPLEX_BODY_TPL, encoding="utf-8") as f:
        body = json.load(f)
    # 模板里 4 组「查询日期」参数全部改为报告截止日
    for p in body.get("dynamicParams", []):
        if p.get("name") == "查询日期":
            p["defaultValue"] = query_date
        elif p.get("name") in ("本期开始",):
            p["defaultValue"] = period_start
        elif p.get("name") in ("本期结束",):
            p["defaultValue"] = period_end
    req = urllib.request.Request(
        f"{BASE}/api/card/{COMPLEX_CDID}/data",
        data=json.dumps(body).encode("utf-8"),
        headers={"Content-Type": "application/json", "Cookie": f"uIdToken={token}"},
        method="POST",
    )
    with urllib.request.urlopen(req, timeout=180) as resp:
        raw = resp.read().decode("utf-8")
    o = json.loads(raw)
    if "Not Login" in raw[:500] and not (o.get("response") or o):
        raise RuntimeError("uIdToken 失效（Not Login），请重新复制浏览器 Cookie 更新 guancli 配置")
    r = o.get("response") or o
    html = (r.get("viewData") or {}).get("html") or ""
    if not html:
        raise RuntimeError(f"复杂报表响应无 html，keys={list(r.keys())}")
    return html


def parse_excel_html(html):
    """把 Aspose.Cell 渲染的 Excel 风格 HTML 解析为二维表。

    复杂报表响应含多个 <table>：完整数据表会出现两次（正常渲染 + 冻结表头副本），
    另有仅 2 行的表头残片。策略：逐 table 解析 → 按内容签名去重 → 拼接，
    既消除重复又保留报表内真实存在的多个不同区块。
    单元格处理：rowspan/colspan 展开占位，数值优先取 x:num 属性（精确值）。
    """
    tables = re.findall(r"<table[^>]*>.*?</table>", html, re.S)
    if not tables:
        tables = [html]
    blocks, sigs = [], []
    for t in tables:
        grid = {}
        row_idx = 0
        for tr in re.finditer(r"<tr[^>]*>(.*?)</tr>", t, re.S):
            col_idx = 0
            for td in re.finditer(r"<td([^>]*)>(.*?)</td>", tr.group(1), re.S):
                attrs, inner = td.group(1), td.group(2)
                while (row_idx, col_idx) in grid:
                    col_idx += 1
                rowspan = int(re.search(r'rowspan="(\d+)"', attrs).group(1)) if "rowspan=" in attrs else 1
                colspan = int(re.search(r'colspan="(\d+)"', attrs).group(1)) if "colspan=" in attrs else 1
                m = re.search(r'x:num="([^"]+)"', attrs)
                if m:
                    try:
                        val = round(float(m.group(1)), 4)
                    except ValueError:
                        val = m.group(1)
                else:
                    txt = re.sub(r"<[^>]+>", "", inner)
                    txt = (txt.replace("&nbsp;", " ").replace("&amp;", "&")
                              .replace("&lt;", "<").replace("&gt;", ">"))
                    val = txt.strip()
                for dr in range(rowspan):
                    for dc in range(colspan):
                        grid[(row_idx + dr, col_idx + dc)] = val if (dr == 0 and dc == 0) else ""
                col_idx += colspan
            row_idx += 1
        if not grid:
            continue
        max_r = max(r for r, _ in grid)
        max_c = max(c for _, c in grid)
        table = [[grid.get((r, c), "") for c in range(max_c + 1)] for r in range(max_r + 1)]
        while table and all(v == "" for v in table[-1]):
            table.pop()
        if not table:
            continue
        sig = json.dumps(table, ensure_ascii=False, default=str)
        if sig in sigs:
            continue  # 重复副本（冻结表头渲染），跳过
        sigs.append(sig)
        blocks.append(table)
    # 过滤「冻结表头残片」：行数 ≤3 且内容等于某个更大区块的头部若干行
    def is_header_fragment(b):
        if len(b) > 3:
            return False
        bs = json.dumps(b, ensure_ascii=False, default=str)
        for other in blocks:
            if other is b or len(other) <= len(b):
                continue
            head = json.dumps(other[:len(b)], ensure_ascii=False, default=str)
            if head == bs:
                return True
        return False
    blocks = [b for b in blocks if not is_header_fragment(b)]
    # 拼接多个不同区块，区块间空一行
    out = []
    for i, b in enumerate(blocks):
        if i:
            out.append([])
        out += b
    return out


# ---------------------------------------------------------------- 主流程
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--date", help="报告截止日 YYYY-MM-DD，缺省=昨日")
    ap.add_argument("--out", help="输出 xlsx 路径")
    args = ap.parse_args()

    if not BASE:
        sys.exit(
            "❌ 缺少环境变量 BI_BASE。请按 README「配置」设置观远 BI 地址，"
            "例如 export BI_BASE=https://bi.example.com"
        )

    end_date = args.date or (dt.date.today() - dt.timedelta(days=1)).isoformat()
    start_date = "2026-01-01"  # 用户规则：本期开始/开始时间固定为 2026-01-01
    out_path = args.out or os.path.join(PROJECT, "data", f"写字楼经营分析-BI取数_{end_date}.xlsx")
    os.makedirs(os.path.dirname(os.path.abspath(out_path)), exist_ok=True)

    param_vals = {"查询日期": end_date, "本期开始": start_date, "本期结束": end_date,
                  "开始时间": start_date, "结束时间": end_date}

    log(f"报告截止日={end_date}　期间开始={start_date}")

    # ---- token 活性预检：失效则自动走 OA→BI SSO 刷新（scripts/bi_auth_refresh.js）----
    refresh_js = os.path.join(SCRIPTS_DIR, "bi_auth_refresh.js")
    node_bin = shutil.which(os.environ.get("NODE_BIN") or "node")
    if os.path.exists(refresh_js) and node_bin:
        env = dict(os.environ)
        if os.environ.get("NODE_PATH"):
            env["NODE_PATH"] = os.environ["NODE_PATH"]
        chk = subprocess.run(
            [node_bin, refresh_js, "--check"],
            capture_output=True, text=True, timeout=60, env=env, shell=False,
        )
        if chk.returncode != 0:
            log("⚠️ 现有 token 已失效，自动刷新（OA→BI SSO）...")
            rf = subprocess.run(
                [node_bin, refresh_js],
                capture_output=True, text=True, timeout=180, env=env, shell=False,
            )
            for line in (rf.stdout + rf.stderr).strip().splitlines():
                log(f"  [refresh] {line}")
            if rf.returncode != 0:
                sys.exit("❌ token 自动刷新失败，请人工检查 OA 凭证或网络后重跑")
            log("✅ token 自动刷新成功")
    elif os.path.exists(refresh_js):
        warn("未找到 node，跳过 token 自动刷新；请设置 NODE_BIN 或安装 Node.js")

    token = get_token()
    log("uIdToken 已读取")

    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    except ImportError:
        sys.exit("❌ 缺 openpyxl，请先安装到 venv")
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    header_fill = PatternFill("solid", fgColor="1A3C6E")
    header_font = Font(color="FFFFFF", bold=True)

    def clean(v):
        """BI 数据可能含 Excel 非法控制字符，写入前清洗"""
        if isinstance(v, str):
            v = ILLEGAL_CHARACTERS_RE.sub("", v)
            if len(v) > 32000:
                v = v[:32000]
            # 千分位纯数字 → 数值型，便于报告端计算（保留前导零的编码类字符串不转）
            s = v.replace(",", "").strip()
            if s and re.fullmatch(r"-?\d+(\.\d+)?", s) and not (len(s) > 1 and s.lstrip("-").startswith("0")):
                try:
                    f = float(s)
                    return int(f) if f.is_integer() else f
                except ValueError:
                    return v
            return v
        return v

    def write_sheet(name, rows, header_row=True):
        ws = wb.create_sheet(name[:31])
        for ri, row in enumerate(rows, 1):
            for ci, v in enumerate(row, 1):
                cell = ws.cell(row=ri, column=ci, value=clean(v))
                if header_row and ri == 1:
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.alignment = Alignment(horizontal="center")
        ws.freeze_panes = "A2"
        return ws

    # ① 复杂报表（B 权责主表）
    b_sheet_name = f"B权责累计分析(至{end_date})"
    try:
        html = fetch_complex_report(token, end_date, start_date, end_date)
        table = parse_excel_html(html)
        if not table:
            raise RuntimeError("HTML 解析结果为空")
        # 期间标签（如「2026年1-9月」）用于命名
        m = re.search(r"2026年1-(\d+)月", json.dumps(table, ensure_ascii=False))
        if m:
            b_sheet_name = f"B权责1-{m.group(1)}月累计分析"
        write_sheet(b_sheet_name, table, header_row=False)
        log(f"✅ {b_sheet_name}: {len(table)} 行 × {len(table[0]) if table else 0} 列（复杂报表）")
        # 关键校验：收入合计不应全 0
        flat = json.dumps(table, ensure_ascii=False)
        if "12," not in flat and not any(isinstance(v, (int, float)) and v > 100 for row in table for v in row):
            warn(f"{b_sheet_name} 疑似收入为空，请核对 Card 3 数据集状态")
    except Exception as e:
        warn(f"B 权责复杂报表取数失败: {e}")
        write_sheet(b_sheet_name, [["取数失败", str(e)]], header_row=False)

    # ② 普通卡（guancli CSV，带分区回退）
    # zero_fallback=True 的卡：分区存在但数值全 0 也视为无数据（合同类 dt 分区滞后）
    ZERO_FALLBACK = {"C合同情况(本年累计)"}
    for sheet, cdid, pnames, limit, critical in CARDS:
        params = {n: param_vals[n] for n in pnames}
        pstr = ",".join(f"{k}={v}" for k, v in params.items()) or "无"
        try:
            rows, used_dt, note = fetch_with_fallback(
                sheet, cdid, pnames, limit, param_vals, end_date,
                zero_fallback=(sheet in ZERO_FALLBACK))
            n = len(rows) - 1 if rows else 0
            if n <= 0 or rows_all_zero(rows) and sheet in ZERO_FALLBACK:
                msg = f"{sheet}（{cdid[:8]}）无有效数据，参数[{pstr}]"
                if critical:
                    warn("关键卡 " + msg + " → sheet 留空，报告须报警，严禁编造")
                else:
                    warn(msg)
                write_sheet(sheet, [[f"⚠️ 无数据（参数: {pstr}）{('；' + note) if note else ''}"]], header_row=False)
            else:
                write_sheet(sheet, rows)
                log(f"✅ {sheet}: {n} 行（参数 {pstr}）{('；' + note) if note else ''}")
        except Exception as e:
            warn(f"{sheet} 取数失败: {e}")
            write_sheet(sheet, [["取数失败", str(e)]], header_row=False)

    # ③ 取数日志 sheet
    log_rows = [["时间", dt.datetime.now().strftime("%Y-%m-%d %H:%M:%S")],
                ["报告截止日", end_date], ["期间开始", start_date],
                ["参数规则", "查询日期=截止日；本期开始/开始时间=2026-01-01；本期结束/结束时间=截止日"],
                [], ["--- 日志 ---"]]
    log_rows += [[l] for l in log_lines]
    if warnings:
        log_rows += [[], ["--- ⚠️ 警告（报告必须引用）---"]]
        log_rows += [[w] for w in warnings]
    write_sheet("取数日志", log_rows, header_row=False)

    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    wb.save(out_path)
    log(f"📦 已保存: {out_path}")
    if warnings:
        log(f"共 {len(warnings)} 条警告，详见『取数日志』sheet")
        sys.exit(2)  # 非零退出码提示有警告，供定时任务判断
    log("全部卡片取数成功，无警告")


if __name__ == "__main__":
    main()
